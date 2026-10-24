"""Buat src/sample.json dari file Excel XBRL IDX.

Pemakaian: python tools/sample.py FinancialStatement-2025-Tahunan-AALI.xlsx > src/sample.json
"""
import json, re, sys
from pathlib import Path
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
info = {}
for r in wb['1000000'].iter_rows(values_only=True):
    if len(r) >= 3 and r[2] and r[1] is not None:
        info[str(r[2]).strip()] = str(r[1])
facts = []
for ws in wb.worksheets:
    m = re.match(r'^\d([235])\d{5}$', ws.title)
    if not m:
        continue
    k = {'2': 'bs', '3': 'is', '5': 'cf'}[m.group(1)]
    for r in ws.iter_rows(min_row=4, values_only=True):
        if len(r) < 4:
            continue
        c, p = r[1], r[2]
        if not isinstance(c, (int, float)) and not isinstance(p, (int, float)):
            continue
        facts.append({'k': k, 's': ws.title, 'id': r[0], 'en': r[3],
                      'c': c if isinstance(c, (int, float)) else None,
                      'p': p if isinstance(p, (int, float)) else None})
print(json.dumps({'source': 'excel', 'file': Path(sys.argv[1]).name, 'info': info, 'facts': facts},
                 ensure_ascii=False, separators=(',', ':')))
